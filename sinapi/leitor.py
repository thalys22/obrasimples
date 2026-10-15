"""Leitura do xlsx oficial do SINAPI. Só Python puro + openpyxl (requirements-etl.txt).

Valida cabeçalhos e falha com ErroLayout se o arquivo mudar, em vez de importar lixo.
"""
import re
from decimal import Decimal

ABAS = {"sd": ("ISD", "CSD"), "cd": ("ICD", "CCD"), "se": ("ISE", "CSE")}
LINHA_CAB = 10
RE_CODIGO_FORMULA = re.compile(r"MATCH\(\s*(\d+)\s*,")


class ErroLayout(Exception):
    pass


def _n(v):
    return " ".join(str(v if v is not None else "").replace("\n", " ").split()).casefold()


def _dec(v):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    try:
        return Decimal(str(v).strip().replace(",", ".")) if isinstance(v, str) else Decimal(repr(float(v))) if isinstance(v, float) else Decimal(v)
    except Exception:
        raise ErroLayout(f"Valor numérico inválido no arquivo: {v!r}")


def abrir(caminho, formulas=False):
    import openpyxl
    return openpyxl.load_workbook(caminho, read_only=True, data_only=not formulas)


def _aba(wb, nome):
    if nome not in wb.sheetnames:
        raise ErroLayout(f"Aba '{nome}' não encontrada. Abas do arquivo: {', '.join(wb.sheetnames)}.")
    return wb[nome]


def _cabecalho(ws, esperado, nome, linha=LINHA_CAB):
    linhas = list(ws.iter_rows(min_row=1, max_row=linha, values_only=True))
    if len(linhas) < linha:
        raise ErroLayout(f"Aba '{nome}': menos de {linha} linhas.")
    cab = linhas[linha - 1]
    for i, texto in enumerate(esperado):
        got = _n(cab[i] if i < len(cab) else None)
        if got != _n(texto):
            raise ErroLayout(f"Aba '{nome}': cabeçalho da coluna {i + 1} esperado '{texto}', veio '{got}'. O layout do arquivo mudou.")
    return linhas


def ler_mes(wb, aba):
    """Mês de referência 'MM/AAAA' da linha 3. Devolve (ano, mes)."""
    ws = _aba(wb, aba)
    for r in ws.iter_rows(min_row=3, max_row=3, values_only=True):
        for v in r[:4]:
            m = re.fullmatch(r"\s*(\d{2})/(\d{4})\s*", str(v or ""))
            if m:
                return int(m.group(2)), int(m.group(1))
    raise ErroLayout(f"Aba '{aba}': mês de referência (linha 3) não encontrado.")


def _colunas_uf(linha, ufs, inicio, nome):
    cols = {}
    for i, v in enumerate(linha):
        t = str(v or "").strip()
        if i >= inicio and t in ufs and t not in cols:
            cols[t] = i
    faltam = set(ufs) - set(cols)
    if faltam:
        raise ErroLayout(f"Aba '{nome}': UF(s) {', '.join(sorted(faltam))} não encontradas no cabeçalho.")
    return cols


def ler_precos(wb, aba, ufs):
    """{codigo: {classificacao, descricao, unidade, coleta, precos:{uf: Decimal|None}}}"""
    ws = _aba(wb, aba)
    linhas = _cabecalho(ws, ["Classificação", "Código do Insumo", "Descrição do Insumo", "Unidade", "Origem de Preço"], aba)
    cols = _colunas_uf(linhas[LINHA_CAB - 1], ufs, 5, aba)
    out = {}
    for r in ws.iter_rows(min_row=LINHA_CAB + 1, values_only=True):
        if r[1] in (None, ""):
            continue
        try:
            cod = int(str(r[1]).strip())
        except ValueError:
            raise ErroLayout(f"Aba '{aba}': código de insumo inválido {r[1]!r}.")
        if cod in out:
            raise ErroLayout(f"Aba '{aba}': código de insumo {cod} repetido.")
        out[cod] = dict(classificacao=str(r[0] or "").strip(), descricao=str(r[2] or "").strip(), unidade=str(r[3] or "").strip(),
                        coleta=str(r[4] or "").strip(), precos={uf: _dec(r[c]) if c < len(r) else None for uf, c in cols.items()})
    return out


def _codigo_formula(v, aba, linha):
    if isinstance(v, (int, float)) and int(v) > 0:
        return int(v)
    m = RE_CODIGO_FORMULA.search(str(v or ""))
    if not m:
        raise ErroLayout(f"Aba '{aba}', linha {linha}: não consegui extrair o código da composição de {str(v)[:60]!r}. A fórmula HYPERLINK mudou?")
    return int(m.group(1))


def ler_custos(wb_formulas, wb_valores, aba, ufs):
    """Duas passadas: fórmulas (código) e valores (custo, %AS). {codigo: {grupo, descricao, unidade, custos:{uf:(custo|None, pct)}}}"""
    wf, wv = _aba(wb_formulas, aba), _aba(wb_valores, aba)
    linhas = _cabecalho(wv, ["Grupo", "Código da Composição", "Descrição", "Unidade"], aba)
    uf_linha = linhas[LINHA_CAB - 2]
    cols = _colunas_uf(uf_linha, ufs, 4, aba)
    sub = linhas[LINHA_CAB - 1]
    for uf, c in cols.items():
        if not (_n(sub[c]).startswith("custo") and _n(sub[c + 1] if c + 1 < len(sub) else None) == "%as"):
            raise ErroLayout(f"Aba '{aba}': colunas de {uf} não são 'Custo (R$)' + '%AS'.")
    out = {}
    it_f = wf.iter_rows(min_row=LINHA_CAB + 1, values_only=True)
    it_v = wv.iter_rows(min_row=LINHA_CAB + 1, values_only=True)
    n = LINHA_CAB
    for rv in it_v:
        n += 1
        rf = next(it_f, None)
        if rf is None:
            raise ErroLayout(f"Aba '{aba}': a passada de fórmulas terminou antes da de valores (linha {n}).")
        if rv[2] in (None, "") and rf[1] in (None, ""):
            continue
        cod = _codigo_formula(rf[1], aba, n)
        if cod in out:
            raise ErroLayout(f"Aba '{aba}': composição {cod} repetida.")
        custos = {}
        for uf, c in cols.items():
            custo = _dec(rv[c]) if c < len(rv) else None
            pct = _dec(rv[c + 1]) if c + 1 < len(rv) else None
            custos[uf] = (custo if custo and custo > 0 else None, pct or Decimal(0))
        out[cod] = dict(grupo=str(rv[0] or "").strip(), descricao=str(rv[2] or "").strip(), unidade=str(rv[3] or "").strip(), custos=custos)
    return out


def ler_analitico(wb):
    """(composicoes, itens, avisos). composicoes: {codigo: {grupo, descricao, unidade, situacao}};
    itens: [(composicao, ordem, tipo_item, codigo_item, descricao, unidade, coeficiente|None, situacao)]"""
    ws = _aba(wb, "Analítico")
    _cabecalho(ws, ["Grupo", "Código da Composição", "Tipo Item", "Código do Item", "Descrição", "Unidade", "Coeficiente", "Situação"], "Analítico")
    comps, itens, avisos = {}, [], []
    ordem = {}
    for n, r in enumerate(ws.iter_rows(min_row=LINHA_CAB + 1, values_only=True), LINHA_CAB + 1):
        if r[1] in (None, ""):
            continue
        cod = int(str(r[1]).strip())
        tipo = _n(r[2])
        if tipo == "":
            if cod in comps:
                raise ErroLayout(f"Analítico, linha {n}: composição {cod} definida duas vezes.")
            comps[cod] = dict(grupo=str(r[0] or "").strip(), descricao=str(r[4] or "").strip(), unidade=str(r[5] or "").strip(), situacao=str(r[7] or "").strip())
            continue
        if tipo not in ("insumo", "composicao"):
            raise ErroLayout(f"Analítico, linha {n}: Tipo Item desconhecido {r[2]!r}.")
        if cod not in comps:
            raise ErroLayout(f"Analítico, linha {n}: item antes do cabeçalho da composição {cod}.")
        coef = _dec(r[6])
        if coef is None or coef == 0:
            avisos.append(f"Coeficiente zero/vazio: composição {cod}, item {r[3]} (linha {n}).")
        ordem[cod] = ordem.get(cod, 0) + 1
        itens.append((cod, ordem[cod], tipo, int(str(r[3]).strip()), str(r[4] or "").strip(), str(r[5] or "").strip(), coef, str(r[7] or "").strip()))
    return comps, itens, avisos
