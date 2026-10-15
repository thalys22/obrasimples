import os
import tempfile
from datetime import date
from decimal import Decimal
from io import StringIO

from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import SimpleTestCase, TestCase

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None

from . import carga, leitor, mapas, servico
from .models import Carga, Composicao, ComposicaoItem, CustoComposicao, Insumo, PrecoInsumo, Referencia

BASE = date(2026, 9, 1)
INTERNO = "Custos Horários Produtivo e Improdutivo dos Equipamentos"


def _formula(c):
    return f'=HYPERLINK("#"&CELL("address",OFFSET(Analítico!$B$1,MATCH({c},Analítico!$B:$B,0)-1,3)),{c})'


def montar_xlsx(caminho, mes="09/2026", ciclo=False, custo12_se=74.10, cab_insumo="Código do Insumo", formula_ruim=False, comp_orfa=False):
    """Planilha pequena no formato do arquivo real: cabeçalho na linha 10, código da composição como fórmula HYPERLINK."""
    wb = openpyxl.Workbook()
    wb.active.title = "Menu"

    def topo(ws, titulo):
        ws.append([titulo])
        ws.append(["RELATÓRIO"])
        ws.append(["Mês de Referência:", mes])
        for _ in range(6):
            ws.append([])

    isd = wb.create_sheet("ISD")
    topo(isd, "SINAPI")
    isd.append(["Classificação", cab_insumo, "Descrição do Insumo", "Unidade", "Origem de\nPreço", "SE", "SP"])
    isd.append(["MATERIAL", "100", "CIMENTO PORTLAND", "KG", "C", 2.00, 1.80])
    isd.append(["MATERIAL", "200", "AREIA MÉDIA", "M3", "CR", None, 100.00])
    isd.append(["MAO DE OBRA", "300", "SERVENTE", "H", "C", 20.00, 22.00])
    isd.append(["MAO DE OBRA", "400", "PEDREIRO", "H", "C", 25.00, 27.00])

    csd = wb.create_sheet("CSD")
    topo(csd, "SINAPI")
    csd.row_dimensions  # noqa
    csd.cell(row=9, column=5, value="SE")
    csd.cell(row=9, column=7, value="SP")
    csd.append(["Grupo", "Código da\nComposição", "Descrição", "Unidade", "Custo (R$)", "%AS", "Custo (R$)", "%AS"])
    codigo = (lambda c: "=HYPERLINK(sem match)") if formula_ruim else _formula
    linhas = [
        ("Paredes", 10, "ARGAMASSA DE ASSENTAMENTO", "M3", 720.00, 0.1667, 660.00, 0),
        ("Paredes", 11, "SERVENTE COM ENCARGOS COMPLEMENTARES", "H", 20.00, 0, 22.00, 0),
        ("Paredes", 12, "ALVENARIA DE VEDAÇÃO", "M2", custo12_se, 0.0971, 77.30, 0),
        ("Paredes", 13, "ALVENARIA DUPLA", "M2", 155.40, 0.0773, 161.20, 0),
        ("Paredes", 14, "ITEM SEM CUSTO", "UN", None, 0, None, 0),
        (INTERNO, 15, "CHP DE BETONEIRA", "CHP", 2.00, 0, 1.80, 0),
    ]
    if ciclo:
        linhas += [("Paredes", 16, "CICLO A", "UN", None, 0, None, 0), ("Paredes", 17, "CICLO B", "UN", None, 0, None, 0)]
    for g, c, d, u, se, pse, sp, psp in linhas:
        csd.append([g, codigo(c), d, u, se, pse, sp, psp])

    an = wb.create_sheet("Analítico")
    topo(an, "SINAPI")
    an.append(["Grupo", "Código da\nComposição", "Tipo Item", "Código do\nItem", "Descrição", "Unidade", "Coeficiente", "Situação"])

    def comp(g, c, d, u, sit, itens):
        an.append([g, str(c), None, None, d, u, None, sit])
        for t, i, desc, un, co, st in itens:
            an.append([g, str(c), t, str(i), desc, un, co, st])

    comp("Paredes", 10, "ARGAMASSA DE ASSENTAMENTO", "M3", "COM CUSTO", [("INSUMO", 100, "CIMENTO PORTLAND", "KG", 300, "COM PREÇO"), ("INSUMO", 200, "AREIA MÉDIA", "M3", 1.2, "COM PREÇO")])
    comp("Paredes", 11, "SERVENTE COM ENCARGOS COMPLEMENTARES", "H", "COM CUSTO", [("INSUMO", 300, "SERVENTE", "H", 1, "COM PREÇO")])
    comp("Paredes", 12, "ALVENARIA DE VEDAÇÃO", "M2", "COM CUSTO", [
        ("COMPOSICAO", 10, "ARGAMASSA DE ASSENTAMENTO", "M3", 0.03, "COM CUSTO"),
        ("COMPOSICAO", 11, "SERVENTE COM ENCARGOS COMPLEMENTARES", "H", 2, "COM CUSTO"),
        ("INSUMO", 400, "PEDREIRO", "H", 0.5, "COM PREÇO")])
    comp("Paredes", 13, "ALVENARIA DUPLA", "M2", "COM CUSTO", [
        ("COMPOSICAO", 12, "ALVENARIA DE VEDAÇÃO", "M2", 2, "COM CUSTO"),
        ("COMPOSICAO", 10, "ARGAMASSA DE ASSENTAMENTO", "M3", 0.01, "COM CUSTO")])
    comp("Paredes", 14, "ITEM SEM CUSTO", "UN", "SEM CUSTO", [("INSUMO", 999, "INSUMO SEM PREÇO NACIONAL", "UN", 1, "SEM PREÇO")])
    comp(INTERNO, 15, "CHP DE BETONEIRA", "CHP", "COM CUSTO", [("INSUMO", 100, "CIMENTO PORTLAND", "KG", 1, "COM PREÇO")])
    if ciclo:
        comp("Paredes", 16, "CICLO A", "UN", "SEM CUSTO", [("COMPOSICAO", 17, "CICLO B", "UN", 1, "SEM CUSTO")])
        comp("Paredes", 17, "CICLO B", "UN", "SEM CUSTO", [("COMPOSICAO", 16, "CICLO A", "UN", 1, "SEM CUSTO")])
    if comp_orfa:
        an.append(["Paredes", "13", "COMPOSICAO", "888", "FANTASMA", "UN", 1, "COM CUSTO"])
    wb.save(caminho)


@__import__("unittest").skipUnless(openpyxl, "openpyxl não instalado (requirements-etl.txt)")
class BaseSinapi(TestCase):
    def setUp(self):
        self.dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.dir.cleanup)
        self.arq = os.path.join(self.dir.name, "sinapi.xlsx")

    def carregar(self, **kw):
        montar_xlsx(self.arq, **kw)
        return carga.carregar(self.arq, BASE, ["SE"], ["sd"], saida=lambda *_: None)


class CargaTests(BaseSinapi):
    def test_carga_le_codigo_da_formula_hyperlink(self):
        self.carregar()
        self.assertEqual(sorted(CustoComposicao.objects.values_list("composicao__codigo", flat=True)), [10, 11, 12, 13, 14, 15])
        c12 = CustoComposicao.objects.get(composicao__codigo=12)
        self.assertEqual(c12.custo_publicado, Decimal("74.10"))
        self.assertEqual(c12.pct_as, Decimal("0.0971"))

    def test_estrutura_guardada_como_veio(self):
        rel = self.carregar()
        self.assertEqual(Composicao.objects.count(), 6)
        self.assertEqual(rel["profundidade_max"], 3)
        self.assertEqual(Composicao.objects.get(codigo=12).itens.count(), 3)
        self.assertEqual(Composicao.objects.get(codigo=12).itens.filter(tipo_item="composicao").count(), 2)

    def test_insumos_sao_a_uniao_de_precos_e_analitico(self):
        self.carregar()
        self.assertEqual(Insumo.objects.count(), 5)
        i = Insumo.objects.get(codigo=999)
        self.assertTrue(i.sem_preco_nacional)
        self.assertEqual(i.classificacao, "")
        self.assertEqual(Insumo.objects.get(codigo=300).tipo, "mao-de-obra")
        self.assertFalse(Insumo.objects.get(codigo=100).sem_preco_nacional)

    def test_fallback_para_sp_marca_origem(self):
        self.carregar()
        ref = Referencia.objects.get(uf="SE")
        areia = PrecoInsumo.objects.get(referencia=ref, insumo__codigo=200)
        self.assertEqual((areia.preco, areia.origem_preco), (Decimal("100"), "sp"))
        cim = PrecoInsumo.objects.get(referencia=ref, insumo__codigo=100)
        self.assertEqual((cim.preco, cim.origem_preco, cim.coleta), (Decimal("2"), "local", "C"))
        self.assertFalse(PrecoInsumo.objects.filter(referencia=ref, insumo__codigo=999).exists())
        self.assertFalse(Referencia.objects.filter(uf="SP").exists(), "SP só serve de fallback, não vira referência sem ser pedida")

    def test_grupo_interno_nao_e_orcavel(self):
        self.carregar()
        self.assertFalse(Composicao.objects.get(codigo=15).orcavel)
        self.assertTrue(Composicao.objects.get(codigo=12).orcavel)

    def test_idempotente_por_sha256(self):
        self.carregar()
        n = (Insumo.objects.count(), ComposicaoItem.objects.count(), PrecoInsumo.objects.count(), Carga.objects.count())
        rel = carga.carregar(self.arq, BASE, ["SE"], ["sd"], saida=lambda *_: None)
        self.assertTrue(rel.get("nada"))
        self.assertEqual(n, (Insumo.objects.count(), ComposicaoItem.objects.count(), PrecoInsumo.objects.count(), Carga.objects.count()))

    def test_segunda_uf_do_mesmo_arquivo_nao_duplica_estrutura(self):
        self.carregar()
        itens = ComposicaoItem.objects.count()
        carga.carregar(self.arq, BASE, ["SE", "SP"], ["sd"], saida=lambda *_: None)
        self.assertEqual(ComposicaoItem.objects.count(), itens)
        self.assertEqual(Referencia.objects.count(), 2)
        sp = PrecoInsumo.objects.get(referencia__uf="SP", insumo__codigo=200)
        self.assertEqual(sp.origem_preco, "local")

    def test_conferencia_aceita_centavo_e_registra_ate_5_por_cento(self):
        rel = self.carregar(custo12_se=74.12)  # 2 centavos de diferença: dentro da tolerância
        self.assertEqual(rel["referencias"]["SE/sd"]["conferencia"], dict(dentro=5, ate5=0, acima5=0, sem_calculo=0))
        Referencia.objects.all().delete()
        Carga.objects.all().delete()
        rel = self.carregar(custo12_se=75.00)  # +1,2%: fora da tolerância, dentro de 5%
        self.assertEqual(rel["referencias"]["SE/sd"]["conferencia"]["ate5"], 1)

    def test_custo_muito_errado_reprova_e_nao_grava_nada(self):
        with self.assertRaisesMessage(carga.ErroCarga, "Conferência de custo reprovada"):
            self.carregar(custo12_se=120.00)
        for m in (Referencia, Insumo, Composicao, ComposicaoItem, PrecoInsumo, CustoComposicao, Carga):
            self.assertEqual(m.objects.count(), 0, m.__name__)

    def test_mes_diferente_do_pedido_falha(self):
        with self.assertRaisesMessage(carga.ErroCarga, "08/2026"):
            self.carregar(mes="08/2026")

    def test_cabecalho_trocado_falha_com_mensagem_clara(self):
        with self.assertRaisesMessage(leitor.ErroLayout, "layout do arquivo mudou"):
            self.carregar(cab_insumo="Cód.")

    def test_formula_de_codigo_irreconhecivel_falha(self):
        with self.assertRaisesMessage(leitor.ErroLayout, "HYPERLINK"):
            self.carregar(formula_ruim=True)

    def test_ciclo_de_auxiliares_falha(self):
        with self.assertRaisesMessage(carga.ErroCarga, "Ciclo"):
            self.carregar(ciclo=True)

    def test_auxiliar_nao_definida_falha(self):
        with self.assertRaisesMessage(carga.ErroCarga, "888"):
            self.carregar(comp_orfa=True)

    def test_uf_inexistente_no_arquivo_falha(self):
        montar_xlsx(self.arq)
        with self.assertRaisesMessage(leitor.ErroLayout, "ZZ"):
            carga.carregar(self.arq, BASE, ["ZZ"], ["sd"], saida=lambda *_: None)

    def test_comando_valida_base_e_roda(self):
        montar_xlsx(self.arq)
        with self.assertRaises(CommandError):
            call_command("carregar_sinapi", arquivo=self.arq, base="2026/09", uf=["SE"], regime=["sd"], stdout=StringIO())
        out = StringIO()
        call_command("carregar_sinapi", arquivo=self.arq, base="2026-09", uf=["SE"], regime=["sd"], stdout=out)
        self.assertIn("SE/sd", out.getvalue())
        self.assertEqual(Referencia.objects.count(), 1)


class ServicoTests(BaseSinapi):
    def setUp(self):
        super().setUp()
        self.carregar()
        self.args = ("SE", "2026-09", "sd")

    def test_referencias_disponiveis(self):
        self.assertEqual(servico.referencias_disponiveis(), [dict(uf="SE", data_base="2026-09", regime="sd")])

    def test_busca_ignora_acento_e_exige_todas_as_palavras(self):
        r = servico.buscar_composicoes("alvenaria vedacao", *self.args)
        self.assertEqual([i["codigo"] for i in r["itens"]], [12])
        self.assertEqual(r["itens"][0]["unidade"], "m²")
        self.assertEqual(servico.buscar_composicoes("ALVENARIA", *self.args)["total"], 2)

    def test_busca_exclui_sem_custo_e_internas_por_padrao(self):
        self.assertEqual(servico.buscar_composicoes("sem custo", *self.args)["total"], 0)
        self.assertEqual(servico.buscar_composicoes("betoneira", *self.args)["total"], 0)
        self.assertEqual(servico.buscar_composicoes("sem custo", *self.args, incluir_sem_custo=True)["total"], 1)
        self.assertEqual(servico.buscar_composicoes("betoneira", *self.args, incluir_internas=True)["total"], 1)

    def test_busca_pagina_e_limita(self):
        a = servico.buscar_composicoes("ar", *self.args, limite=2)
        self.assertEqual((a["total"], len(a["itens"])), (4, 2))  # ARGAMASSA, SERVENTE COM ENCARGOS... contêm "ar"
        b = servico.buscar_composicoes("ar", *self.args, limite=2, deslocamento=2)
        self.assertEqual(len(b["itens"]), 2)
        self.assertFalse({i["codigo"] for i in a["itens"]} & {i["codigo"] for i in b["itens"]})
        self.assertEqual(len(servico.buscar_composicoes("ar", *self.args, limite=9999)["itens"]), 4)

    def test_busca_curta_e_referencia_inexistente(self):
        with self.assertRaises(servico.SinapiErro):
            servico.buscar_composicoes("a", *self.args)
        with self.assertRaises(servico.SinapiErro):
            servico.buscar_composicoes("alvenaria", "BA", "2026-09", "sd")
        with self.assertRaises(servico.SinapiErro):
            servico.buscar_composicoes("alvenaria", "SE", "2026-08", "sd")

    def test_obter_achata_somando_insumo_por_dois_caminhos(self):
        c = servico.obter_composicao(13, *self.args)
        por = {i["codigo_insumo"]: i for i in c["itens"]}
        # cimento: 2*0,03*300 + 0,01*300 = 21 ; areia: 2*0,03*1,2 + 0,01*1,2 = 0,084
        self.assertEqual(por[100]["coef"], Decimal("21"))
        self.assertEqual(por[200]["coef"], Decimal("0.084"))
        self.assertEqual(por[300]["coef"], Decimal("4"))
        self.assertEqual(por[400]["coef"], Decimal("1"))
        self.assertEqual(c["custo_calculado"], Decimal("155.40"))
        self.assertEqual(c["custo"], Decimal("155.4000"))
        self.assertTrue(any(l["tipo"] == "composicao" and l["nivel"] == 1 for l in c["estrutura"]))

    def test_obter_marca_preco_de_sp_e_avisa(self):
        c = servico.obter_composicao(12, *self.args)
        areia = next(i for i in c["itens"] if i["codigo_insumo"] == 200)
        self.assertEqual(areia["origem_preco"], "sp")
        self.assertTrue(any("São Paulo" in a and "SE" in a for a in c["avisos"]))
        self.assertEqual(c["custo_calculado"], Decimal("74.10"))

    def test_obter_preco_ausente_e_none_nunca_zero(self):
        c = servico.obter_composicao(14, *self.args)
        self.assertIsNone(c["itens"][0]["preco"])
        self.assertIsNone(c["custo_calculado"])
        self.assertTrue(any("sem preço" in a for a in c["avisos"]))
        self.assertTrue(any("não publicou custo" in a for a in c["avisos"]))

    def test_obter_inexistente(self):
        with self.assertRaises(servico.SinapiErro):
            servico.obter_composicao(777, *self.args)


class MapasTests(SimpleTestCase):
    def test_unidades_e_tipos(self):
        self.assertEqual(mapas.unidade_prumo("M2"), "m²")
        self.assertEqual(mapas.unidade_prumo("m3"), "m³")
        self.assertEqual(mapas.unidade_prumo("M3XKM"), "m3xkm")
        self.assertTrue(mapas.unidade_conhecida("M2xMES"))
        self.assertFalse(mapas.unidade_conhecida("XYZ"))
        self.assertEqual(mapas.tipo_do_insumo("MAO DE OBRA"), "mao-de-obra")
        self.assertEqual(mapas.tipo_do_insumo("EQUIPAMENTO (AQUISIÇÃO)"), "equipamento")
        self.assertEqual(mapas.tipo_do_insumo("SERVIÇOS"), "empreitada")

    def test_tolerancia(self):
        self.assertEqual(carga.tolerancia(0.01), 0.02)
        self.assertAlmostEqual(carga.tolerancia(1000), 5.0)
